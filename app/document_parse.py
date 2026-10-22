"""Document parsing for the personal data center (phase 3, iteration 3.1).

Bytes in -> plain text out. Content is validated by magic bytes rather than
the uploaded extension alone. Nothing here executes file contents, writes to
disk, or calls the network.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB upload cap
MAX_TEXT_CHARS = 200_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}

_PDF_MAGIC = b"%PDF-"
_DOCX_MAGIC = b"PK\x03\x04"
_ZIP_ENTRY_XML = b"[Content_Types].xml"
# Markers used to detect UTF-16 text encodings.
_UTF16_BOMS = (b"\xff\xfe", b"\xfe\xff")


class DocumentParseError(ValueError):
    """Raised when an upload is too large, of a bad type, corrupt, or empty."""


@dataclass(frozen=True)
class Page:
    """One extracted page/section. page_hint is 1-based for PDF, None otherwise."""

    text: str
    page_hint: int | None = None
    heading: str | None = None


def safe_filename(filename: str) -> str:
    """Reduce an uploaded name to a safe basename (no path traversal)."""
    name = (filename or "").replace("\\", "/").split("/")[-1].strip()
    name = re.sub(r"[\x00-\x1f]", "", name)
    name = re.sub(r"[<>:\"/\\|?*]", "_", name)
    name = name.lstrip(".") or "未命名文件"
    return name[:255]


def normalize_extension(filename: str) -> str:
    lower = filename.lower()
    for ext in ALLOWED_EXTENSIONS:
        if lower.endswith(ext):
            return ".docx" if ext == ".docx" else ext
    return ""


def validate_extension(filename: str) -> str:
    ext = normalize_extension(filename)
    if not ext:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise DocumentParseError(f"不支持的文件类型，仅允许：{allowed}")
    return ext


def _decode_plain_text(data: bytes) -> str:
    if data.startswith(_UTF16_BOMS):
        for encoding in ("utf-16", "utf-16-le", "utf-16-be"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("无法识别文本文件编码（请使用 UTF-8）")


def _looks_like_text(data: bytes) -> bool:
    if data.startswith(_UTF16_BOMS):
        return True
    if b"\x00" in data[:4096]:
        return False
    try:
        data[:4096].decode("utf-8")
        return True
    except UnicodeDecodeError:
        try:
            data[:4096].decode("gb18030")
            return True
        except UnicodeDecodeError:
            return False


def _parse_pdf(data: bytes) -> list[Page]:
    if not data.startswith(_PDF_MAGIC):
        raise DocumentParseError("文件扩展名为 .pdf，但内容不是有效的 PDF")
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentParseError("服务器缺少 PDF 解析依赖 pypdf") from exc
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        pages: list[Page] = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            text = text.strip()
            if text:
                pages.append(Page(text=text, page_hint=index))
        return pages
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"PDF 已损坏或无法解析：{type(exc).__name__}") from exc


def _parse_docx(data: bytes) -> list[Page]:
    if not data.startswith(_DOCX_MAGIC):
        raise DocumentParseError("文件扩展名为 .docx，但内容不是有效的 DOCX（zip）")
    if _ZIP_ENTRY_XML not in data:
        raise DocumentParseError("文件扩展名为 .docx，但缺少 Office 文档结构")
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("服务器缺少 DOCX 解析依赖 python-docx") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"DOCX 已损坏或无法解析：{type(exc).__name__}") from exc

    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    joined = "\n".join(blocks).strip()
    return [Page(text=joined)] if joined else []


def parse_document(filename: str, data: bytes) -> list[Page]:
    """Validate and parse an uploaded file into ordered text pages."""
    if not data:
        raise DocumentParseError("文件内容为空")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentParseError(f"文件超过大小上限 {MAX_FILE_BYTES // (1024 * 1024)}MB")
    safe_name = safe_filename(filename)
    ext = validate_extension(safe_name)

    if ext == ".pdf":
        pages = _parse_pdf(data)
    elif ext == ".docx":
        pages = _parse_docx(data)
    else:
        if not _looks_like_text(data):
            raise DocumentParseError("文本文件包含无法解码的二进制内容")
        text = _decode_plain_text(data).strip()
        pages = [Page(text=text)] if text else []

    total = sum(len(p.text) for p in pages)
    if total == 0:
        raise DocumentParseError("未能从文件中提取到任何文本（可能是扫描件/空文档）")
    if total > MAX_TEXT_CHARS:
        raise DocumentParseError(f"解析后文本超过 {MAX_TEXT_CHARS} 字符上限")
    return pages
