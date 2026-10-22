"""Phase 3 iterations 3.1-3.3: file upload, parsing, chunking, chunk retrieval."""

from __future__ import annotations

import io
import sqlite3

import pytest
from fastapi.testclient import TestClient

from app import document_chunk, document_parse, retriever
from app.document_parse import DocumentParseError, Page
from app.main import app, init_db

client = TestClient(app)


@pytest.fixture
def isolated_db(tmp_path, monkeypatch):
    db_path = tmp_path / "phase3.db"
    monkeypatch.setenv("DATABASE_PATH", str(db_path))
    monkeypatch.setattr("app.main.DB_PATH", db_path)
    init_db()
    return db_path


# ---------- minimal valid file builders ----------

def _minimal_pdf(text: str = "Coffee preference is americano") -> bytes:
    # Build a tiny single-page PDF with a correct xref table.
    objs: list[bytes] = []
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    objs.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>")
    stream = b"BT /F1 14 Tf 72 720 Td (" + text.encode("latin-1", "replace") + b") Tj ET"
    objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n"
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode() +
            b" /Root 1 0 R >>\nstartxref\n" + str(xref_pos).encode() + b"\n%%EOF")
    return bytes(out)


def _minimal_docx(paragraphs: list[str]) -> bytes:
    import docx
    document = docx.Document()
    for para in paragraphs:
        document.add_paragraph(para)
    document.add_heading("项目决策", level=1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "技术栈"
    table.rows[0].cells[1].text = "FastAPI SQLite"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ---------- parser unit tests ----------

def test_safe_filename_blocks_traversal():
    assert "/" not in document_parse.safe_filename("../../etc/passwd")
    assert document_parse.safe_filename("..\\..\\a:b?.txt").endswith(".txt")


def test_parse_txt_and_md_basic():
    pages = document_parse.parse_document("note.txt", "你好世界\n第二行".encode())
    assert len(pages) == 1 and "你好世界" in pages[0].text

    pages_md = document_parse.parse_document("a.md", "# 标题\n\n正文内容".encode())
    assert pages_md and "正文内容" in pages_md[0].text


def test_parse_pdf_real():
    pages = document_parse.parse_document("a.pdf", _minimal_pdf())
    assert pages and pages[0].page_hint == 1
    assert "americano" in pages[0].text.lower()


def test_parse_docx_real():
    pages = document_parse.parse_document("a.docx", _minimal_docx(["第一段项目背景", "第二段技术选择"]))
    text = "\n".join(p.text for p in pages)
    assert "第一段项目背景" in text and "FastAPI SQLite" in text


def test_reject_bad_extension_and_masquerading_file():
    with pytest.raises(DocumentParseError):
        document_parse.parse_document("a.exe", b"MZ\x90\x00")
    with pytest.raises(DocumentParseError):
        document_parse.parse_document("a.pdf", b"this is plain text, not a pdf")
    with pytest.raises(DocumentParseError):
        document_parse.parse_document("a.docx", b"definitely not a zip")


def test_reject_corrupt_pdf_and_docx():
    with pytest.raises(DocumentParseError):
        document_parse.parse_document("a.pdf", b"%PDF-broken junk without structure")
    # Valid zip magic but not an office document
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(DocumentParseError):
        document_parse.parse_document("a.docx", buf.getvalue())


def test_reject_empty_and_oversized():
    with pytest.raises(DocumentParseError):
        document_parse.parse_document("a.txt", b"")
    big = b"a" * (document_parse.MAX_FILE_BYTES + 1)
    with pytest.raises(DocumentParseError):
        document_parse.parse_document("a.txt", big)


# ---------- chunking + provenance ----------

def test_chunk_markdown_tracks_heading_and_provenance():
    md = "# 架构设计\n\n" + "系统采用 FastAPI 与 SQLite。" * 40 + "\n\n## 检索模块\n\n关键词检索是默认实现。" * 5
    chunks = document_chunk.chunk_pages([Page(text=md)], ext=".md")
    assert chunks, "expected chunks"
    assert all(c.char_count > 0 for c in chunks)
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    first = chunks[0]
    assert first.heading == "架构设计"
    later = [c for c in chunks if c.heading == "检索模块"]
    assert later, "heading 检索模块 not tracked"


def test_chunk_pdf_keeps_page_hint():
    pages = [Page(text="第一页内容 " * 100, page_hint=1), Page(text="第二页内容 " * 100, page_hint=2)]
    chunks = document_chunk.chunk_pages(pages, ext=".pdf")
    page_hints = {c.page_hint for c in chunks}
    assert 1 in page_hints and 2 in page_hints


# ---------- API: text + upload, dedup, delete cascade ----------

def test_text_endpoint_creates_chunks(isolated_db):
    resp = client.post("/api/documents/text", json={
        "filename": "学习笔记.md", "content": "# 标题\n\nFastAPI 是一个现代 Python 框架。" * 30,
    })
    assert resp.status_code == 200
    doc_id = resp.json()["id"]
    detail = client.get(f"/api/documents/{doc_id}").json()
    assert detail["status"] == "parsed" and detail["chunk_count"] >= 1
    assert detail["chunks"] and detail["chunks"][0]["document_id" if False else "chunk_index"] >= 0
    assert all(c["excerpt"] for c in detail["chunks"])


def test_upload_txt_and_duplicate_rejected(isolated_db):
    content = ("个人项目的技术决策记录：使用 SQLite 作为第一版存储。" * 10).encode("utf-8")
    r1 = client.post("/api/documents/upload",
                     files={"file": ("决策记录.txt", content, "text/plain")})
    assert r1.status_code == 200, r1.text
    assert r1.json()["source_type"] == "upload" and r1.json()["chunk_count"] >= 1

    r2 = client.post("/api/documents/upload",
                     files={"file": ("另一个名字.txt", content, "text/plain")})
    assert r2.status_code == 409


def test_upload_docx_roundtrip(isolated_db):
    data = _minimal_docx(["项目使用 FastAPI 后端", "待解决问题是检索排序"])
    resp = client.post("/api/documents/upload",
                       files={"file": ("项目说明.docx", data, "application/vnd.openxmlformats")})
    assert resp.status_code == 200, resp.text
    doc_id = resp.json()["id"]
    detail = client.get(f"/api/documents/{doc_id}").json()
    joined = " ".join(c["excerpt"] for c in detail["chunks"])
    assert "FastAPI" in joined


def test_upload_rejects_unsupported_and_corrupt(isolated_db):
    bad_ext = client.post("/api/documents/upload", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert bad_ext.status_code == 400
    fake_pdf = client.post("/api/documents/upload", files={"file": ("x.pdf", b"not a pdf", "application/pdf")})
    assert fake_pdf.status_code == 400
    # failure must leave no residue
    assert client.get("/api/documents").json() == []


def test_chunk_search_ranks_and_delete_cascades(isolated_db):
    client.post("/api/documents/text", json={
        "filename": "咖啡.md", "content": "我的咖啡偏好是喝美式咖啡，不加糖。" * 5,
    })
    other = client.post("/api/documents/text", json={
        "filename": "运动.md", "content": "每周跑步三次，力量训练两次。" * 5,
    })
    with sqlite3.connect(isolated_db) as conn:
        conn.row_factory = sqlite3.Row
        hits = retriever.search_chunks(conn, "咖啡偏好 美式")
        assert hits, "expected chunk hits"
        assert "咖啡" in hits[0]["filename"]
        assert hits[0]["score"] > 0
        coffee_doc = hits[0]["document_id"]

    # deleting the document removes its chunks from retrieval
    assert client.delete(f"/api/documents/{coffee_doc}").status_code == 200
    with sqlite3.connect(isolated_db) as conn:
        conn.row_factory = sqlite3.Row
        hits_after = retriever.search_chunks(conn, "咖啡偏好 美式")
        assert all(h["document_id"] != coffee_doc for h in hits_after)
        leftover = conn.execute("SELECT COUNT(*) FROM document_chunks WHERE document_id=?", (coffee_doc,)).fetchone()[0]
        assert leftover == 0
    # unrelated doc survives
    assert client.get(f"/api/documents/{other.json()['id']}").status_code == 200


def test_chat_uses_chunk_level_citations(isolated_db, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    client.post("/api/documents/text", json={
        "filename": "目标.md", "content": "下个月的学习目标是掌握 RAG 检索与重排序。" * 5,
    })
    resp = client.post("/api/chat", json={"message": "我下个月的学习目标是什么？RAG 检索"})
    assert resp.status_code == 200
    used = resp.json()["used_documents"]
    assert used and {"document_id", "chunk_id", "filename", "score"} <= set(used[0])
    assert "chunk_id" in used[0]


def test_model_unavailable_does_not_block_document_management(isolated_db, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    r = client.post("/api/documents/text", json={"filename": "x.txt", "content": "离线也能管理的资料内容 " * 10})
    assert r.status_code == 200
    assert client.get("/api/documents").status_code == 200
    assert client.delete(f"/api/documents/{r.json()['id']}").json()["deleted"] is True


# ---------- legacy migration ----------

def test_legacy_documents_migrate_and_become_searchable(tmp_path, monkeypatch):
    db_path = tmp_path / "legacy3.db"
    monkeypatch.setenv("DATABASE_PATH", str(db_path))
    monkeypatch.setattr("app.main.DB_PATH", db_path)

    import sqlite3 as _sql
    with _sql.connect(db_path) as conn:
        # phase-1 style documents table (only 4 columns) plus dependencies
        conn.executescript("""
        CREATE TABLE documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT NOT NULL,
            content TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        INSERT INTO documents(filename,content,created_at)
        VALUES ('旧笔记.txt', '旧资料里记录了 RAG 混合检索的计划', 't0');
        CREATE TABLE memories (
            id INTEGER PRIMARY KEY AUTOINCREMENT, content TEXT NOT NULL,
            category TEXT DEFAULT 'general', source TEXT DEFAULT 'manual',
            confidence REAL DEFAULT 1, confirmed INTEGER DEFAULT 1,
            created_at TEXT, updated_at TEXT
        );
        CREATE TABLE profile (id INTEGER PRIMARY KEY CHECK(id=1), display_name TEXT, goals TEXT, preferences TEXT, updated_at TEXT);
        INSERT INTO profile VALUES(1,'我','','','t0');
        CREATE TABLE conversations (id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT, created_at TEXT);
        CREATE TABLE messages (id INTEGER PRIMARY KEY AUTOINCREMENT, conversation_id INTEGER, role TEXT, content TEXT, created_at TEXT);
        CREATE TABLE approvals (id INTEGER PRIMARY KEY AUTOINCREMENT, action_type TEXT, payload TEXT, status TEXT, created_at TEXT, reviewed_at TEXT);
        """)

    init_db()  # must add columns, chunks table, and backfill chunks

    with _sql.connect(db_path) as conn:
        conn.row_factory = _sql.Row
        doc = conn.execute("SELECT * FROM documents WHERE filename='旧笔记.txt'").fetchone()
        assert doc["status"] == "parsed" and doc["source_type"] == "manual_text"
        chunk_count = conn.execute(
            "SELECT COUNT(*) FROM document_chunks WHERE document_id=?", (doc["id"],)
        ).fetchone()[0]
        assert chunk_count >= 1
        hits = retriever.search_chunks(conn, "RAG 混合检索")
        assert hits and any(h["document_id"] == doc["id"] for h in hits)

    # migration is idempotent
    init_db()
